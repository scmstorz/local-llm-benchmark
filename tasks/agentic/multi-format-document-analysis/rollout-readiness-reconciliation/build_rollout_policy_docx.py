#!/usr/bin/env python3
"""Build the synthetic rollout-policy DOCX fixture."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
GATE_COLUMN_WIDTHS_DXA = [1760, 2720, 1760, 3120]

CALIBRI = "Calibri"
BLACK = "000000"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "666666"
LIGHT_GRAY = "F2F4F7"
BORDER = "B7C9D6"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_style_font(style, *, name: str, size: float, color: str, bold=False, italic=False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    run_properties = style.element.get_or_add_rPr()
    run_fonts = ensure_child(run_properties, "w:rFonts")
    run_fonts.set(qn("w:ascii"), name)
    run_fonts.set(qn("w:hAnsi"), name)
    run_fonts.set(qn("w:eastAsia"), name)


def configure_paragraph_style(
    style,
    *,
    name: str = CALIBRI,
    size: float = 11,
    color: str = BLACK,
    bold=False,
    italic=False,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.0,
    keep_with_next=False,
    keep_together=False,
) -> None:
    set_style_font(
        style,
        name=name,
        size=size,
        color=color,
        bold=bold,
        italic=italic,
    )
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    paragraph.line_spacing = line_spacing
    paragraph.keep_with_next = keep_with_next
    paragraph.keep_together = keep_together
    paragraph.widow_control = True


def configure_styles(document: Document) -> None:
    styles = document.styles
    configure_paragraph_style(
        styles["Normal"], size=11, after=6, line_spacing=1.10
    )
    configure_paragraph_style(
        styles["Heading 1"],
        size=16,
        color=BLUE,
        bold=True,
        before=16,
        after=8,
        line_spacing=1.0,
        keep_with_next=True,
        keep_together=True,
    )
    configure_paragraph_style(
        styles["Heading 2"],
        size=13,
        color=BLUE,
        bold=True,
        before=12,
        after=6,
        line_spacing=1.0,
        keep_with_next=True,
        keep_together=True,
    )
    configure_paragraph_style(
        styles["Heading 3"],
        size=12,
        color=DARK_BLUE,
        bold=True,
        before=8,
        after=4,
        line_spacing=1.0,
        keep_with_next=True,
        keep_together=True,
    )

    masthead_title = styles.add_style("Masthead Title", WD_STYLE_TYPE.PARAGRAPH)
    configure_paragraph_style(
        masthead_title,
        size=23,
        bold=True,
        after=4,
        line_spacing=1.0,
        keep_with_next=True,
    )
    masthead_subtitle = styles.add_style(
        "Masthead Subtitle", WD_STYLE_TYPE.PARAGRAPH
    )
    configure_paragraph_style(
        masthead_subtitle,
        size=14,
        color="373737",
        after=12,
        line_spacing=1.0,
        keep_with_next=True,
    )
    metadata = styles.add_style("Masthead Metadata", WD_STYLE_TYPE.PARAGRAPH)
    configure_paragraph_style(metadata, size=10.5, after=2, line_spacing=1.0)
    table_citation = styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    configure_paragraph_style(
        table_citation,
        size=9,
        color=MUTED,
        italic=True,
        before=4,
        after=4,
        line_spacing=1.0,
        keep_with_next=True,
    )
    table_header = styles.add_style("Gate Table Header", WD_STYLE_TYPE.PARAGRAPH)
    configure_paragraph_style(
        table_header,
        size=9.5,
        color=DARK_BLUE,
        bold=True,
        after=0,
        line_spacing=1.0,
        keep_together=True,
    )
    table_body = styles.add_style("Gate Table Body", WD_STYLE_TYPE.PARAGRAPH)
    configure_paragraph_style(
        table_body,
        size=9.5,
        after=0,
        line_spacing=1.0,
        keep_together=True,
    )
    numbered = styles.add_style("Policy Number", WD_STYLE_TYPE.PARAGRAPH)
    configure_paragraph_style(
        numbered,
        size=11,
        after=8,
        line_spacing=1.167,
        keep_together=True,
    )

    metadata_label = styles.add_style("Metadata Label", WD_STYLE_TYPE.CHARACTER)
    set_style_font(
        metadata_label,
        name=CALIBRI,
        size=10.5,
        color=BLACK,
        bold=True,
    )
    directive_label = styles.add_style("Directive Label", WD_STYLE_TYPE.CHARACTER)
    set_style_font(
        directive_label,
        name=CALIBRI,
        size=11,
        color=DARK_BLUE,
        bold=True,
    )

    configure_paragraph_style(
        styles["Header"], size=8.5, color=MUTED, after=0, line_spacing=1.0
    )
    configure_paragraph_style(
        styles["Footer"], size=8.5, color=MUTED, after=0, line_spacing=1.0
    )


def add_numbering_definition(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    alignment = OxmlElement("w:lvlJc")
    alignment.set(qn("w:val"), "left")
    level.extend([start, number_format, level_text, alignment])

    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "720")
    indentation.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_properties.extend([tabs, indentation, spacing])
    level.append(paragraph_properties)
    abstract.append(level)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    num_properties = ensure_child(properties, "w:numPr")
    level = ensure_child(num_properties, "w:ilvl")
    level.set(qn("w:val"), "0")
    number = ensure_child(num_properties, "w:numId")
    number.set(qn("w:val"), str(num_id))


def add_bookmarked_heading(
    document: Document,
    text: str,
    bookmark_name: str,
    bookmark_id: int,
):
    paragraph = document.add_paragraph(style="Heading 1")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    paragraph._p.append(start)
    paragraph.add_run(text)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(end)
    return paragraph


def set_cell_margins(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = ensure_child(cell_properties, "w:tcMar")
    for side, width in CELL_MARGINS_DXA.items():
        margin = ensure_child(margins, f"w:{side}")
        margin.set(qn("w:w"), str(width))
        margin.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    width_element = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcW")
    width_element.set(qn("w:type"), "dxa")
    width_element.set(qn("w:w"), str(width))
    cell.width = Twips(width)


def set_cell_fill(cell, color: str) -> None:
    shading = ensure_child(cell._tc.get_or_add_tcPr(), "w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)


def set_table_borders(table) -> None:
    borders = ensure_child(table._tbl.tblPr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = ensure_child(borders, f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), BORDER)


def set_table_geometry(table) -> None:
    if sum(GATE_COLUMN_WIDTHS_DXA) != CONTENT_WIDTH_DXA:
        raise AssertionError("Gate table columns must sum to 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_properties = table._tbl.tblPr
    table_width = ensure_child(table_properties, "w:tblW")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    table_indent = ensure_child(table_properties, "w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = ensure_child(table_properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    caption = ensure_child(table_properties, "w:tblCaption")
    caption.set(qn("w:val"), "Regional-gates")
    description = ensure_child(table_properties, "w:tblDescription")
    description.set(
        qn("w:val"),
        "Four mandatory region-level production rollout gates",
    )

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in GATE_COLUMN_WIDTHS_DXA:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, GATE_COLUMN_WIDTHS_DXA[column_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    set_table_borders(table)


def set_paragraph_border_bottom(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = ensure_child(properties, "w:pBdr")
    bottom = ensure_child(borders, "w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), DARK_BLUE)
    paragraph.paragraph_format.space_after = Pt(10)


def append_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, text, separate, result, end):
        paragraph.add_run()._r.append(element)


def configure_section(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.style = document.styles["Header"]
    header.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5), WD_TAB_ALIGNMENT.RIGHT
    )
    header.add_run("Rollout Governance Memo")
    header.add_run("\tAtlas regional launch")

    footer = section.footer.paragraphs[0]
    footer.style = document.styles["Footer"]
    footer.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5), WD_TAB_ALIGNMENT.RIGHT
    )
    footer.add_run("Synthetic benchmark fixture")
    footer.add_run("\tPage ")
    append_field(footer, " PAGE ")
    footer.add_run(" of ")
    append_field(footer, " NUMPAGES ")


def add_metadata(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph(style="Masthead Metadata")
    label_run = paragraph.add_run(f"{label}: ")
    label_run.style = document.styles["Metadata Label"]
    paragraph.add_run(value)


def set_table_cell_text(cell, text: str, *, header: bool) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.style = "Gate Table Header" if header else "Gate Table Body"
    paragraph.add_run(text)
    if header:
        set_cell_fill(cell, LIGHT_GRAY)


def build_document() -> Document:
    document = Document()
    configure_styles(document)
    configure_section(document)
    numbering_id = add_numbering_definition(document)

    properties = document.core_properties
    properties.author = "Local LLM Benchmark Lab"
    properties.last_modified_by = "Local LLM Benchmark Lab"
    properties.title = "Regional Rollout Governance Policy"
    properties.subject = "Synthetic multi-format benchmark fixture"
    properties.keywords = "synthetic benchmark, rollout policy, regional gates"
    properties.comments = (
        "Project-authored synthetic policy fixture for read-only document analysis."
    )
    fixed_timestamp = datetime(2026, 8, 31, 18, 0, tzinfo=timezone.utc)
    properties.created = fixed_timestamp
    properties.modified = fixed_timestamp

    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings.element.append(update_fields)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    document.add_paragraph("ROLLOUT GOVERNANCE MEMO", style="Masthead Title")
    document.add_paragraph(
        "Regional production-release authorization for the Atlas service rollout",
        style="Masthead Subtitle",
    )
    add_metadata(document, "To", "Regional Launch Steering Committee")
    add_metadata(document, "From", "Rollout Governance Office")
    add_metadata(document, "Date", "31 August 2026")
    add_metadata(
        document,
        "Re",
        "Region-level readiness gates and phased release authority",
    )
    add_metadata(document, "Status", "Approved policy for cutoff review")
    rule = document.add_paragraph()
    set_paragraph_border_bottom(rule)

    directive = document.add_paragraph()
    directive_label = directive.add_run("Directive. ")
    directive_label.style = document.styles["Directive Label"]
    directive.add_run(
        "Evaluate launch eligibility per region. Portfolio summaries are context, not authorization."
    )

    add_bookmarked_heading(document, "Regional gates", "Regional_gates", 1)
    document.add_paragraph(
        "A region passes only when every criterion below is satisfied for that same region."
    )
    document.add_paragraph("Table 1. Regional gates", style="Table Citation")

    gate_rows = [
        ("Gate", "Passing criterion", "Evaluation", "Evidence basis"),
        (
            "Validation completion",
            "Completion is at least 95%.",
            "Per region",
            "Final validation evidence at the cutoff",
        ),
        (
            "Open Sev-1 issues",
            "Open Sev-1 issue count equals 0.",
            "Per region",
            "Final reliability review at the cutoff",
        ),
        (
            "Open Sev-2 issues",
            "Open Sev-2 issue count is at most 2.",
            "Per region",
            "Final reliability review at the cutoff",
        ),
        (
            "Support capacity",
            "Support capacity / projected peak load is at least 1.25.",
            "Per region",
            "Regional readiness workbook, Capacity sheet",
        ),
    ]
    table = document.add_table(rows=len(gate_rows), cols=4)
    for row_index, row_values in enumerate(gate_rows):
        for column_index, value in enumerate(row_values):
            set_table_cell_text(
                table.cell(row_index, column_index),
                value,
                header=row_index == 0,
            )
    set_table_geometry(table)
    document.add_paragraph(
        "Policy owner: Rollout Governance Office. Applies to the 2026-08-31 cutoff review.",
        style="Table Citation",
    )

    decision_rule_heading = add_bookmarked_heading(
        document, "Decision rule", "Decision_rule", 2
    )
    decision_rule_heading.paragraph_format.page_break_before = True
    document.add_paragraph(
        "Apply every gate independently to each region. A region is eligible only when all four gates pass for that region. Portfolio averages, rollups or dashboard status cannot replace or override a failed regional gate."
    )

    procedure_heading = document.add_paragraph(
        "Decision procedure", style="Heading 2"
    )
    procedure_heading.paragraph_format.keep_with_next = True
    for item in (
        "Evaluate each region against all four criteria using evidence finalized by the cutoff.",
        "Authorize only regions that pass every criterion; do not transfer strength between regions.",
        "Keep each failing region on hold until its failed gate is cleared and revalidated.",
    ):
        paragraph = document.add_paragraph(item, style="Policy Number")
        apply_numbering(paragraph, numbering_id)

    add_bookmarked_heading(document, "Phased release", "Phased_release", 3)
    document.add_paragraph(
        "A phased production launch may proceed when at least one region passes all four gates. Only regions that pass every gate may launch; all other regions remain on hold until their failed gates are cleared and the evidence is revalidated."
    )

    add_bookmarked_heading(document, "Evidence cutoff", "Evidence_cutoff", 4)
    document.add_paragraph(
        "Use evidence finalized no later than 2026-08-31 18:00 UTC. Later changes belong to a new decision cycle and may not be used retroactively for this review."
    )
    return document


def main() -> None:
    args = parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(args.output)


if __name__ == "__main__":
    main()
