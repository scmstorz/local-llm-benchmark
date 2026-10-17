#!/usr/bin/env python3
"""Run deterministic structural and semantic checks for the DOCX fixture."""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "r": R_NS}
W = f"{{{W_NS}}}"

EXPECTED_TABLE = [
    ["Gate", "Passing criterion", "Evaluation", "Evidence basis"],
    [
        "Validation completion",
        "Completion is at least 95%.",
        "Per region",
        "Final validation evidence at the cutoff",
    ],
    [
        "Open Sev-1 issues",
        "Open Sev-1 issue count equals 0.",
        "Per region",
        "Final reliability review at the cutoff",
    ],
    [
        "Open Sev-2 issues",
        "Open Sev-2 issue count is at most 2.",
        "Per region",
        "Final reliability review at the cutoff",
    ],
    [
        "Support capacity",
        "Support capacity / projected peak load is at least 1.25.",
        "Per region",
        "Regional readiness workbook, Capacity sheet",
    ],
]
EXPECTED_HEADINGS = [
    "Regional gates",
    "Decision rule",
    "Decision procedure",
    "Phased release",
    "Evidence cutoff",
]
EXPECTED_BOOKMARKS = {
    "Regional_gates",
    "Decision_rule",
    "Phased_release",
    "Evidence_cutoff",
}
EXPECTED_GRID = [1760, 2720, 1760, 3120]
EXPECTED_CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--document", required=True, type=Path)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def integer_attribute(element, name: str) -> int:
    if element is None:
        raise AssertionError(f"Missing element for attribute {name}")
    value = element.get(f"{W}{name}")
    if value is None:
        raise AssertionError(f"Missing w:{name} on {element.tag}")
    return int(value)


def style_element(styles_root, style_id: str):
    matches = styles_root.xpath(
        f'./w:style[@w:styleId="{style_id}"]', namespaces=NS
    )
    if len(matches) != 1:
        raise AssertionError(f"Expected one style {style_id}, found {len(matches)}")
    return matches[0]


def assert_style(
    styles_root,
    style_id: str,
    *,
    font: str,
    half_points: int,
    before: int,
    after: int,
    line: int,
    color: str | None = None,
) -> None:
    style = style_element(styles_root, style_id)
    fonts = style.find("w:rPr/w:rFonts", namespaces=NS)
    if fonts is None or fonts.get(f"{W}ascii") != font or fonts.get(f"{W}hAnsi") != font:
        raise AssertionError(f"{style_id}: unexpected fonts")
    size = style.find("w:rPr/w:sz", namespaces=NS)
    if integer_attribute(size, "val") != half_points:
        raise AssertionError(f"{style_id}: unexpected font size")
    spacing = style.find("w:pPr/w:spacing", namespaces=NS)
    observed = {
        "before": integer_attribute(spacing, "before"),
        "after": integer_attribute(spacing, "after"),
        "line": integer_attribute(spacing, "line"),
    }
    expected = {"before": before, "after": after, "line": line}
    if observed != expected or spacing.get(f"{W}lineRule") != "auto":
        raise AssertionError(
            f"{style_id}: expected spacing {expected!r}, found {observed!r}"
        )
    if color is not None:
        color_element = style.find("w:rPr/w:color", namespaces=NS)
        if color_element is None or color_element.get(f"{W}val") != color:
            raise AssertionError(f"{style_id}: unexpected color")


def main() -> None:
    args = parse_args()
    document_path = args.document.resolve()
    with zipfile.ZipFile(document_path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"Corrupt ZIP member: {bad_member}")
        members = set(archive.namelist())
        if "word/vbaProject.bin" in members:
            raise AssertionError("Document contains a VBA project")
        if any(name.startswith("word/comments") for name in members):
            raise AssertionError("Document contains comments")
        for relationship_path in sorted(
            name for name in members if name.endswith(".rels")
        ):
            relationship_root = etree.fromstring(archive.read(relationship_path))
            external = relationship_root.xpath(
                './rel:Relationship[@TargetMode="External"]',
                namespaces={"rel": PKG_REL_NS},
            )
            if external:
                raise AssertionError(
                    f"External relationship in {relationship_path}: {len(external)}"
                )
        document_root = etree.fromstring(archive.read("word/document.xml"))
        styles_root = etree.fromstring(archive.read("word/styles.xml"))
        numbering_root = etree.fromstring(archive.read("word/numbering.xml"))
        footer_root = etree.fromstring(archive.read("word/footer1.xml"))

    document = Document(document_path)
    if len(document.sections) != 1:
        raise AssertionError(f"Expected one section, found {len(document.sections)}")
    section = document.sections[0]
    observed_section = {
        "page_width": section.page_width.twips,
        "page_height": section.page_height.twips,
        "top_margin": section.top_margin.twips,
        "right_margin": section.right_margin.twips,
        "bottom_margin": section.bottom_margin.twips,
        "left_margin": section.left_margin.twips,
        "header_distance": section.header_distance.twips,
        "footer_distance": section.footer_distance.twips,
    }
    expected_section = {
        "page_width": 12240,
        "page_height": 15840,
        "top_margin": 1440,
        "right_margin": 1440,
        "bottom_margin": 1440,
        "left_margin": 1440,
        "header_distance": 708,
        "footer_distance": 708,
    }
    if observed_section != expected_section:
        raise AssertionError(
            f"Unexpected section geometry: {json.dumps(observed_section, sort_keys=True)}"
        )

    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    if headings != EXPECTED_HEADINGS:
        raise AssertionError(f"Unexpected heading sequence: {headings!r}")
    heading_styles = {
        paragraph.text: paragraph.style.name
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    }
    if heading_styles["Decision procedure"] != "Heading 2":
        raise AssertionError("Decision procedure must use Heading 2")
    if any(
        heading_styles[heading] != "Heading 1"
        for heading in EXPECTED_HEADINGS
        if heading != "Decision procedure"
    ):
        raise AssertionError("Source-anchor headings must use Heading 1")

    bookmarks = {
        item.get(f"{W}name")
        for item in document_root.xpath(".//w:bookmarkStart", namespaces=NS)
    }
    if not EXPECTED_BOOKMARKS.issubset(bookmarks):
        raise AssertionError(f"Missing bookmarks: {EXPECTED_BOOKMARKS - bookmarks}")

    if len(document.tables) != 1:
        raise AssertionError(f"Expected one table, found {len(document.tables)}")
    table_values = [
        [cell.text for cell in row.cells] for row in document.tables[0].rows
    ]
    if table_values != EXPECTED_TABLE:
        raise AssertionError(f"Unexpected Regional-gates table: {table_values!r}")

    table_xml = document_root.find(".//w:tbl", namespaces=NS)
    if table_xml is None:
        raise AssertionError("Regional-gates table is missing")
    table_width = table_xml.find("w:tblPr/w:tblW", namespaces=NS)
    table_indent = table_xml.find("w:tblPr/w:tblInd", namespaces=NS)
    table_caption = table_xml.find("w:tblPr/w:tblCaption", namespaces=NS)
    if integer_attribute(table_width, "w") != 9360 or table_width.get(f"{W}type") != "dxa":
        raise AssertionError("Table width must be 9360 DXA")
    if integer_attribute(table_indent, "w") != 120 or table_indent.get(f"{W}type") != "dxa":
        raise AssertionError("Table indent must be 120 DXA")
    if table_caption is None or table_caption.get(f"{W}val") != "Regional-gates":
        raise AssertionError("Table caption anchor must be Regional-gates")

    grid = [
        integer_attribute(column, "w")
        for column in table_xml.xpath("w:tblGrid/w:gridCol", namespaces=NS)
    ]
    if grid != EXPECTED_GRID:
        raise AssertionError(f"Unexpected table grid: {grid!r}")
    for row_number, row in enumerate(
        table_xml.xpath("w:tr", namespaces=NS), start=1
    ):
        cell_widths = [
            integer_attribute(cell.find("w:tcPr/w:tcW", namespaces=NS), "w")
            for cell in row.xpath("w:tc", namespaces=NS)
        ]
        if cell_widths != EXPECTED_GRID:
            raise AssertionError(
                f"Row {row_number}: cell widths {cell_widths!r} do not match grid"
            )
        for cell in row.xpath("w:tc", namespaces=NS):
            for side, expected_width in EXPECTED_CELL_MARGINS.items():
                margin = cell.find(f"w:tcPr/w:tcMar/w:{side}", namespaces=NS)
                if integer_attribute(margin, "w") != expected_width:
                    raise AssertionError(
                        f"Row {row_number}: unexpected {side} cell margin"
                    )

    header_fills = [
        shading.get(f"{W}fill")
        for shading in table_xml.xpath("w:tr[1]/w:tc/w:tcPr/w:shd", namespaces=NS)
    ]
    if header_fills != ["F2F4F7"] * 4:
        raise AssertionError(f"Unexpected table header fills: {header_fills!r}")

    assert_style(
        styles_root,
        "Normal",
        font="Calibri",
        half_points=22,
        before=0,
        after=120,
        line=264,
        color="000000",
    )
    assert_style(
        styles_root,
        "Heading1",
        font="Calibri",
        half_points=32,
        before=320,
        after=160,
        line=240,
        color="2E74B5",
    )
    assert_style(
        styles_root,
        "Heading2",
        font="Calibri",
        half_points=26,
        before=240,
        after=120,
        line=240,
        color="2E74B5",
    )
    assert_style(
        styles_root,
        "Heading3",
        font="Calibri",
        half_points=24,
        before=160,
        after=80,
        line=240,
        color="1F4D78",
    )
    assert_style(
        styles_root,
        "TableCitation",
        font="Calibri",
        half_points=18,
        before=80,
        after=80,
        line=240,
        color="666666",
    )
    assert_style(
        styles_root,
        "PolicyNumber",
        font="Calibri",
        half_points=22,
        before=0,
        after=160,
        line=280,
        color="000000",
    )

    numbered_paragraphs = document_root.xpath(
        './/w:p[w:pPr/w:pStyle[@w:val="PolicyNumber"]]', namespaces=NS
    )
    if len(numbered_paragraphs) != 3:
        raise AssertionError(
            f"Expected three numbered paragraphs, found {len(numbered_paragraphs)}"
        )
    num_ids = {
        paragraph.find("w:pPr/w:numPr/w:numId", namespaces=NS).get(f"{W}val")
        for paragraph in numbered_paragraphs
    }
    if len(num_ids) != 1:
        raise AssertionError(f"Numbered paragraphs use multiple numIds: {num_ids!r}")
    num_id = next(iter(num_ids))
    number = numbering_root.xpath(
        f'./w:num[@w:numId="{num_id}"]', namespaces=NS
    )
    if len(number) != 1:
        raise AssertionError(f"Missing numbering definition for numId {num_id}")
    abstract_id = number[0].find("w:abstractNumId", namespaces=NS).get(f"{W}val")
    abstract = numbering_root.xpath(
        f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]', namespaces=NS
    )[0]
    number_indent = abstract.find("w:lvl/w:pPr/w:ind", namespaces=NS)
    number_tab = abstract.find("w:lvl/w:pPr/w:tabs/w:tab", namespaces=NS)
    number_spacing = abstract.find("w:lvl/w:pPr/w:spacing", namespaces=NS)
    if (
        integer_attribute(number_indent, "left") != 720
        or integer_attribute(number_indent, "hanging") != 360
        or integer_attribute(number_tab, "pos") != 720
        or integer_attribute(number_spacing, "after") != 160
        or integer_attribute(number_spacing, "line") != 280
    ):
        raise AssertionError("Numbering geometry does not match the preset")

    required_fragments = [
        "A region passes only when every criterion below is satisfied for that same region.",
        "Portfolio averages, rollups or dashboard status cannot replace or override a failed regional gate.",
        "A phased production launch may proceed when at least one region passes all four gates.",
        "Use evidence finalized no later than 2026-08-31 18:00 UTC.",
    ]
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for fragment in required_fragments:
        if fragment not in body_text:
            raise AssertionError(f"Required policy fragment is missing: {fragment}")

    footer_instructions = {
        text.strip()
        for text in footer_root.xpath(".//w:instrText/text()", namespaces=NS)
    }
    if footer_instructions != {"PAGE", "NUMPAGES"}:
        raise AssertionError(
            f"Unexpected footer field instructions: {footer_instructions!r}"
        )
    if document.core_properties.author != "Local LLM Benchmark Lab":
        raise AssertionError("Unexpected document author metadata")

    result = {
        "status": "passed",
        "document": str(document_path),
        "sha256": sha256_file(document_path),
        "size_bytes": document_path.stat().st_size,
        "sections": 1,
        "section_geometry_twips": observed_section,
        "headings": headings,
        "bookmarks": sorted(EXPECTED_BOOKMARKS),
        "table": {
            "caption": "Regional-gates",
            "rows": len(EXPECTED_TABLE),
            "columns": 4,
            "width_dxa": 9360,
            "indent_dxa": 120,
            "grid_dxa": EXPECTED_GRID,
            "cell_margins_dxa": EXPECTED_CELL_MARGINS,
        },
        "numbered_paragraphs": 3,
        "external_relationships": 0,
        "comments": 0,
        "vba_projects": 0,
        "policy_fragments": required_fragments,
        "footer_fields": sorted(footer_instructions),
        "author": document.core_properties.author,
    }
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
