#!/usr/bin/env python3
"""Build the synthetic regional-readiness XLSX fixture."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "17365D"
TEAL = "0F6B78"
GREEN = "2E7D32"
GREEN_LIGHT = "E2F0D9"
AMBER = "C65911"
AMBER_LIGHT = "FCE4D6"
BLUE_LIGHT = "D9EAF7"
GRAY_LIGHT = "F2F2F2"
GRAY = "666666"
WHITE = "FFFFFF"
BODY_FONT = "Liberation Sans"
THIN_GRAY = Side(style="thin", color="B7C9D6")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def set_text_style(cell, *, size=10, bold=False, color="000000") -> None:
    cell.font = Font(name=BODY_FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(vertical="center")


def configure_page(sheet, print_area: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A3"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = print_area
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.oddFooter.center.text = "Synthetic benchmark fixture - regional readiness"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.oddFooter.center.size = 8
    sheet.oddFooter.center.color = GRAY
    sheet.oddFooter.right.size = 8
    sheet.oddFooter.right.color = GRAY


def style_title(sheet, title: str, subtitle: str, last_column: str) -> None:
    sheet.merge_cells(f"A1:{last_column}1")
    title_cell = sheet["A1"]
    title_cell.value = title
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(name=BODY_FONT, size=18, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(f"A2:{last_column}2")
    subtitle_cell = sheet["A2"]
    subtitle_cell.value = subtitle
    subtitle_cell.fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    subtitle_cell.font = Font(name=BODY_FONT, size=10, italic=True, color=NAVY)
    subtitle_cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 28


def build_summary(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "Summary"
    configure_page(sheet, "A1:F11")
    style_title(
        sheet,
        "Regional Readiness Dashboard",
        "Management portfolio view - validate launch eligibility against the regional rollout policy.",
        "F",
    )

    sheet.column_dimensions["A"].width = 29
    sheet.column_dimensions["B"].width = 19
    sheet.column_dimensions["C"].width = 4
    sheet.column_dimensions["D"].width = 3
    sheet.column_dimensions["E"].width = 27
    sheet.column_dimensions["F"].width = 17

    labels = {
        "A3": "Portfolio status",
        "A4": "Average completion",
        "A5": "Average open issues / region",
        "A6": "Portfolio capacity ratio",
    }
    for address, value in labels.items():
        cell = sheet[address]
        cell.value = value
        cell.fill = PatternFill("solid", fgColor=GRAY_LIGHT)
        cell.border = Border(bottom=THIN_GRAY)
        set_text_style(cell, bold=True, color=NAVY)

    sheet["B3"] = '=IF(AND(B4>=$F$4,B5<=$F$5,B6>=$F$6),"GO","HOLD")'
    sheet["B4"] = 0.9625
    sheet["B5"] = 1.25
    sheet["B6"] = "=SUM('Capacity'!B4:B7)/SUM('Capacity'!C4:C7)"
    sheet["B3"].comment = Comment(
        "Intentionally uses portfolio aggregates. This business rule is evidence and must not be repaired.",
        "User",
    )
    sheet["B6"].comment = Comment(
        "Portfolio support capacity divided by portfolio projected peak load.",
        "User",
    )

    for row in range(3, 7):
        cell = sheet[f"B{row}"]
        cell.border = Border(bottom=THIN_GRAY)
        cell.alignment = Alignment(vertical="center", horizontal="center")
        set_text_style(cell, size=11, bold=row == 3, color=NAVY)
    sheet["B3"].font = Font(name=BODY_FONT, size=16, bold=True, color=GREEN)
    sheet["B3"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
    sheet["B4"].number_format = "0.00%"
    sheet["B5"].number_format = "0.00"
    sheet["B6"].number_format = "0.0000"

    sheet.merge_cells("E3:F3")
    sheet["E3"] = "Portfolio thresholds"
    sheet["E3"].fill = PatternFill("solid", fgColor=TEAL)
    sheet["E3"].font = Font(name=BODY_FONT, size=11, bold=True, color=WHITE)
    sheet["E3"].alignment = Alignment(horizontal="left", vertical="center")
    threshold_rows = [
        (4, "Average completion", 0.95, "0%"),
        (5, "Average open issues", 2, "0.00"),
        (6, "Aggregate capacity", "='Capacity'!H3", "0.00"),
    ]
    for row, label, value, number_format in threshold_rows:
        sheet[f"E{row}"] = label
        sheet[f"F{row}"] = value
        sheet[f"E{row}"].fill = PatternFill("solid", fgColor=GRAY_LIGHT)
        sheet[f"F{row}"].fill = PatternFill("solid", fgColor=GRAY_LIGHT)
        sheet[f"E{row}"].border = Border(bottom=THIN_GRAY)
        sheet[f"F{row}"].border = Border(bottom=THIN_GRAY)
        set_text_style(sheet[f"E{row}"], bold=True, color=NAVY)
        set_text_style(sheet[f"F{row}"], color=NAVY)
        sheet[f"F{row}"].alignment = Alignment(horizontal="center", vertical="center")
        sheet[f"F{row}"].number_format = number_format

    sheet.merge_cells("A8:F8")
    sheet["A8"] = "Decision-use note"
    sheet["A8"].fill = PatternFill("solid", fgColor=AMBER)
    sheet["A8"].font = Font(name=BODY_FONT, size=11, bold=True, color=WHITE)
    sheet["A8"].alignment = Alignment(vertical="center")

    sheet.merge_cells("A9:F10")
    sheet["A9"] = (
        "This dashboard applies a portfolio-level aggregate formula. Regional launch authorization "
        "is governed by rollout-policy.docx and must be reconciled with the final reliability review."
    )
    sheet["A9"].fill = PatternFill("solid", fgColor=AMBER_LIGHT)
    sheet["A9"].font = Font(name=BODY_FONT, size=10, color="7F2F00")
    sheet["A9"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A9"].border = Border(left=THIN_GRAY, right=THIN_GRAY, bottom=THIN_GRAY)
    sheet.row_dimensions[9].height = 28
    sheet.row_dimensions[10].height = 18

    sheet.merge_cells("A11:F11")
    sheet["A11"] = "Source basis: portfolio rollup supplied for the 2026-08-31 steering review."
    sheet["A11"].font = Font(name=BODY_FONT, size=8, italic=True, color=GRAY)
    sheet["A11"].alignment = Alignment(vertical="center")

    sheet.conditional_formatting.add(
        "B3",
        FormulaRule(formula=['B3="GO"'], fill=PatternFill("solid", fgColor=GREEN_LIGHT)),
    )


def build_capacity(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Capacity")
    configure_page(sheet, "A1:H10")
    style_title(
        sheet,
        "Regional Support Capacity",
        "Typed capacity inputs with auditable region-level ratios.",
        "H",
    )

    widths = {
        "A": 17,
        "B": 22,
        "C": 24,
        "D": 18,
        "E": 16,
        "F": 4,
        "G": 26,
        "H": 15,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    headers = ["Region", "Support capacity", "Projected peak load", "Capacity ratio", "Gate check"]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name=BODY_FONT, size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 30

    rows = [
        ("North", 137, 108),
        ("Central", 150, 117),
        ("West", 150, 100),
        ("East", 180, 130),
    ]
    for row_number, (region, capacity, peak_load) in enumerate(rows, start=4):
        sheet.cell(row=row_number, column=1, value=region)
        sheet.cell(row=row_number, column=2, value=capacity)
        sheet.cell(row=row_number, column=3, value=peak_load)
        sheet.cell(row=row_number, column=4, value=f"=IFERROR(B{row_number}/C{row_number},0)")
        sheet.cell(row=row_number, column=5, value=f'=IF(D{row_number}>=$H$3,"PASS","HOLD")')
        for column in range(1, 6):
            cell = sheet.cell(row=row_number, column=column)
            set_text_style(cell, color=NAVY)
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "center",
                vertical="center",
            )
        sheet.cell(row=row_number, column=2).number_format = "#,##0"
        sheet.cell(row=row_number, column=3).number_format = "#,##0"
        sheet.cell(row=row_number, column=4).number_format = "0.0000"
        sheet.cell(row=row_number, column=4).comment = Comment(
            "Support capacity divided by projected peak load.",
            "User",
        )

    table = Table(displayName="RegionalCapacityTable", ref="A3:E7")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    sheet["G3"] = "Regional capacity gate"
    sheet["H3"] = 1.25
    sheet["G3"].fill = PatternFill("solid", fgColor=GRAY_LIGHT)
    sheet["H3"].fill = PatternFill("solid", fgColor=GRAY_LIGHT)
    sheet["G3"].border = Border(bottom=THIN_GRAY)
    sheet["H3"].border = Border(bottom=THIN_GRAY)
    set_text_style(sheet["G3"], bold=True, color=NAVY)
    set_text_style(sheet["H3"], bold=True, color=NAVY)
    sheet["H3"].number_format = "0.00"
    sheet["H3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["H3"].comment = Comment(
        "Minimum support-capacity-to-projected-peak-load ratio from the rollout policy.",
        "User",
    )

    sheet.conditional_formatting.add(
        "D4:D7",
        CellIsRule(
            operator="greaterThanOrEqual",
            formula=["$H$3"],
            fill=PatternFill("solid", fgColor=GREEN_LIGHT),
        ),
    )
    sheet.conditional_formatting.add(
        "D4:D7",
        CellIsRule(
            operator="lessThan",
            formula=["$H$3"],
            fill=PatternFill("solid", fgColor=AMBER_LIGHT),
        ),
    )

    sheet.merge_cells("A9:H9")
    sheet["A9"] = "Formula note"
    sheet["A9"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A9"].font = Font(name=BODY_FONT, size=11, bold=True, color=WHITE)
    sheet["A9"].alignment = Alignment(vertical="center")

    sheet.merge_cells("A10:H10")
    sheet["A10"] = (
        "Capacity ratio = support capacity / projected peak load. "
        "Gate check compares each regional ratio with the 1.25 policy threshold."
    )
    sheet["A10"].fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    sheet["A10"].font = Font(name=BODY_FONT, size=10, color=NAVY)
    sheet["A10"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[10].height = 28


def main() -> None:
    args = parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.properties.creator = "Local LLM Benchmark Lab"
    workbook.properties.lastModifiedBy = "Local LLM Benchmark Lab"
    workbook.properties.title = "Regional Readiness"
    workbook.properties.subject = "Synthetic multi-format benchmark fixture"
    workbook.properties.description = (
        "Capacity inputs, formula-derived regional ratios and an intentionally misleading portfolio dashboard."
    )
    fixed_timestamp = datetime(2026, 9, 10, 12, 0, tzinfo=timezone.utc)
    workbook.properties.created = fixed_timestamp
    workbook.properties.modified = fixed_timestamp
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    build_summary(workbook)
    build_capacity(workbook)
    workbook.active = 0
    workbook.save(args.output)


if __name__ == "__main__":
    main()
