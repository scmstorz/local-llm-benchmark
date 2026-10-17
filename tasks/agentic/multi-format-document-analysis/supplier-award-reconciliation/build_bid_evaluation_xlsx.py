#!/usr/bin/env python3
"""Build the synthetic supplier bid-evaluation XLSX fixture."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "17365D"
TEAL = "0F6B78"
GREEN = "2E7D32"
GREEN_LIGHT = "E2F0D9"
AMBER = "C65911"
AMBER_LIGHT = "FCE4D6"
BLUE_LIGHT = "D9EAF7"
GRAY_LIGHT = "F2F2F2"
GRAY = "666666"
WHITE = "FFFFFF"
BODY_FONT = "Liberation Sans"
THIN_GRAY = Side(style="thin", color="B7C9D6")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def set_text_style(cell, *, size=10, bold=False, color="000000") -> None:
    cell.font = Font(name=BODY_FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(vertical="center")


def configure_page(sheet, print_area: str, freeze_panes: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = freeze_panes
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = print_area
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.oddFooter.center.text = "Synthetic benchmark fixture - supplier award"
    sheet.oddFooter.right.text = "Page &P of &N"
    for item in (sheet.oddFooter.center, sheet.oddFooter.right):
        item.size = 8
        item.color = GRAY


def style_title(sheet, title: str, subtitle: str, last_column: str) -> None:
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = title
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].font = Font(name=BODY_FONT, size=18, bold=True, color=WHITE)
    sheet["A1"].alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = subtitle
    sheet["A2"].fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    sheet["A2"].font = Font(name=BODY_FONT, size=10, italic=True, color=NAVY)
    sheet["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 28


def build_dashboard(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "Dashboard"
    configure_page(sheet, "A1:F11", "A3")
    style_title(
        sheet,
        "Commercial Bid Dashboard",
        "Adjusted first-year cost ranking - commercial view only.",
        "F",
    )

    widths = {"A": 30, "B": 34, "C": 4, "D": 4, "E": 25, "F": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    labels = {
        "A3": "Recommended supplier",
        "A4": "Lowest adjusted first-year cost",
        "A5": "Recommendation basis",
    }
    for address, value in labels.items():
        cell = sheet[address]
        cell.value = value
        cell.fill = PatternFill("solid", fgColor=GRAY_LIGHT)
        cell.border = Border(bottom=THIN_GRAY)
        set_text_style(cell, bold=True, color=NAVY)

    sheet["B3"] = "=INDEX('Bid Analysis'!A5:A8,MATCH(MIN('Bid Analysis'!E5:E8),'Bid Analysis'!E5:E8,0))"
    sheet["B4"] = "=MIN('Bid Analysis'!E5:E8)"
    sheet["B5"] = "Lowest adjusted first-year cost"
    sheet["B3"].comment = Comment(
        "Intentionally ranks cost only and does not include assurance eligibility. This rule is benchmark evidence and must not be repaired.",
        "User",
    )
    for row in range(3, 6):
        cell = sheet[f"B{row}"]
        cell.border = Border(bottom=THIN_GRAY)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        set_text_style(cell, size=11, bold=row == 3, color=NAVY)
    sheet["B3"].font = Font(name=BODY_FONT, size=16, bold=True, color=GREEN)
    sheet["B3"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
    sheet["B4"].number_format = '[$EUR ]#,##0'
    sheet.row_dimensions[5].height = 28

    sheet.merge_cells("A7:A8")
    sheet["A7"] = "Decision-use scope"
    sheet["A7"].fill = PatternFill("solid", fgColor=AMBER)
    sheet["A7"].font = Font(name=BODY_FONT, size=11, bold=True, color=WHITE)
    sheet["A7"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet["B7"] = "Commercial ranking only"
    sheet["B8"] = "Assurance findings are not incorporated"
    for address in ("B7", "B8"):
        sheet[address].fill = PatternFill("solid", fgColor=AMBER_LIGHT)
        sheet[address].font = Font(name=BODY_FONT, size=10, color="7F2F00")
        sheet[address].alignment = Alignment(vertical="center", wrap_text=True)
        sheet[address].border = Border(bottom=THIN_GRAY)

    sheet.merge_cells("A10:F10")
    sheet["A10"] = "Governance note"
    sheet["A10"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A10"].font = Font(name=BODY_FONT, size=11, bold=True, color=WHITE)
    sheet.merge_cells("A11:F11")
    sheet["A11"] = (
        "Final award authority remains with supplier-selection-policy.docx and the signed-final "
        "supplier assurance review. This dashboard cannot waive eligibility gates."
    )
    sheet["A11"].fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    sheet["A11"].font = Font(name=BODY_FONT, size=10, color=NAVY)
    sheet["A11"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[11].height = 30
    sheet.conditional_formatting.add(
        "B3",
        FormulaRule(formula=['B3="Cedar"'], fill=PatternFill("solid", fgColor=GREEN_LIGHT)),
    )


def build_bid_analysis(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Bid Analysis")
    configure_page(sheet, "A1:M11", "A5")
    style_title(
        sheet,
        "Supplier Bid Analysis",
        "Typed commercial and capacity inputs with formula-derived eligibility metrics.",
        "M",
    )
    widths = {
        "A": 14, "B": 17, "C": 16, "D": 16, "E": 19,
        "F": 15, "G": 15, "H": 17, "I": 13, "J": 14,
        "K": 3, "L": 24, "M": 16,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    headers = [
        "Supplier", "Base annual fee", "Mobilization", "Transition credit",
        "Adjusted first-year cost", "Available units", "Required units",
        "Capacity headroom", "Cost gate", "Capacity gate",
    ]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name=BODY_FONT, size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[4].height = 34

    suppliers = [
        ("Atlas", 1080000, 120000, 60000, 116, 100),
        ("Boreal", 1120000, 80000, 35000, 125, 100),
        ("Cedar", 1030000, 90000, 40000, 120, 100),
        ("Delta", 1020000, 120000, 10000, 112, 100),
    ]
    for row_number, (supplier, base_fee, mobilization, credit, available, required) in enumerate(suppliers, start=5):
        values = [supplier, base_fee, mobilization, credit, None, available, required]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column, value=value)
        sheet.cell(row=row_number, column=5, value=f"=B{row_number}+C{row_number}-D{row_number}")
        sheet.cell(row=row_number, column=8, value=f"=IFERROR(F{row_number}/G{row_number}-1,0)")
        sheet.cell(row=row_number, column=9, value=f'=IF(E{row_number}<=$M$3,"PASS","HOLD")')
        sheet.cell(row=row_number, column=10, value=f'=IF(H{row_number}>=$M$4,"PASS","HOLD")')
        for column in range(1, 11):
            cell = sheet.cell(row=row_number, column=column)
            set_text_style(cell, color=NAVY)
            cell.alignment = Alignment(horizontal="left" if column == 1 else "center", vertical="center")
        for column in (2, 3, 4, 5):
            sheet.cell(row=row_number, column=column).number_format = '[$EUR ]#,##0'
        sheet.cell(row=row_number, column=8).number_format = "0%"

    table = Table(displayName="SupplierBidTable", ref="A4:J8")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)

    sheet["L3"] = "Adjusted-cost ceiling"
    sheet["M3"] = 1200000
    sheet["L4"] = "Minimum headroom"
    sheet["M4"] = 0.15
    for row in (3, 4):
        for column in (12, 13):
            cell = sheet.cell(row=row, column=column)
            cell.fill = PatternFill("solid", fgColor=GRAY_LIGHT)
            cell.border = Border(bottom=THIN_GRAY)
            set_text_style(cell, bold=column == 12, color=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet["M3"].number_format = '[$EUR ]#,##0'
    sheet["M4"].number_format = "0%"
    sheet["M3"].comment = Comment("Maximum adjusted first-year cost from the selection policy.", "User")
    sheet["M4"].comment = Comment("Minimum capacity headroom from the selection policy.", "User")

    sheet.conditional_formatting.add(
        "H5:H8", CellIsRule(operator="greaterThanOrEqual", formula=["$M$4"], fill=PatternFill("solid", fgColor=GREEN_LIGHT))
    )
    sheet.conditional_formatting.add(
        "H5:H8", CellIsRule(operator="lessThan", formula=["$M$4"], fill=PatternFill("solid", fgColor=AMBER_LIGHT))
    )
    sheet.merge_cells("A10:M10")
    sheet["A10"] = "Formula definitions"
    sheet["A10"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A10"].font = Font(name=BODY_FONT, size=11, bold=True, color=WHITE)
    sheet.merge_cells("A11:M11")
    sheet["A11"] = (
        "Adjusted first-year cost = base annual fee + mobilization - transition credit. "
        "Capacity headroom = available units / required units - 1."
    )
    sheet["A11"].fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    sheet["A11"].font = Font(name=BODY_FONT, size=10, color=NAVY)
    sheet["A11"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[11].height = 28


def main() -> None:
    args = parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.properties.creator = "Local LLM Benchmark Lab"
    workbook.properties.lastModifiedBy = "Local LLM Benchmark Lab"
    workbook.properties.title = "Supplier Bid Evaluation"
    workbook.properties.subject = "Synthetic multi-format benchmark fixture"
    workbook.properties.description = "Supplier bid inputs, derived costs and capacity headroom, plus an intentionally assurance-blind commercial dashboard."
    fixed_timestamp = datetime(2026, 9, 8, 16, 0, tzinfo=timezone.utc)
    workbook.properties.created = fixed_timestamp
    workbook.properties.modified = fixed_timestamp
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    build_dashboard(workbook)
    build_bid_analysis(workbook)
    workbook.active = 0
    workbook.save(args.output)


if __name__ == "__main__":
    main()
