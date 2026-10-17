#!/usr/bin/env python3
"""Build the synthetic supplier-selection-policy DOCX fixture."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
GATE_COLUMN_WIDTHS_DXA = [1800, 2720, 1760, 3080]
CALIBRI = "Calibri"
BLACK = "000000"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "666666"
LIGHT_GRAY = "F2F4F7"
BORDER = "B7C9D6"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_style_font(style, *, name: str, size: float, color: str, bold=False, italic=False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    run_properties = style.element.get_or_add_rPr()
    run_fonts = ensure_child(run_properties, "w:rFonts")
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        run_fonts.set(qn(key), name)


def configure_paragraph_style(
    style, *, size=11, color=BLACK, bold=False, italic=False,
    before=0, after=0, line_spacing=1.0, keep_with_next=False, keep_together=False,
) -> None:
    set_style_font(style, name=CALIBRI, size=size, color=color, bold=bold, italic=italic)
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    paragraph.line_spacing = line_spacing
    paragraph.keep_with_next = keep_with_next
    paragraph.keep_together = keep_together
    paragraph.widow_control = True


def configure_styles(document: Document) -> None:
    styles = document.styles
    configure_paragraph_style(styles["Normal"], size=11, after=6, line_spacing=1.10)
    configure_paragraph_style(styles["Heading 1"], size=16, color=BLUE, bold=True, before=16, after=8, keep_with_next=True, keep_together=True)
    configure_paragraph_style(styles["Heading 2"], size=13, color=BLUE, bold=True, before=12, after=6, keep_with_next=True, keep_together=True)
    configure_paragraph_style(styles["Heading 3"], size=12, color=DARK_BLUE, bold=True, before=8, after=4, keep_with_next=True, keep_together=True)
    additions = [
        ("Masthead Title", WD_STYLE_TYPE.PARAGRAPH, dict(size=23, bold=True, after=4, keep_with_next=True)),
        ("Masthead Subtitle", WD_STYLE_TYPE.PARAGRAPH, dict(size=14, color="373737", after=12, keep_with_next=True)),
        ("Masthead Metadata", WD_STYLE_TYPE.PARAGRAPH, dict(size=10.5, after=2)),
        ("Table Citation", WD_STYLE_TYPE.PARAGRAPH, dict(size=9, color=MUTED, italic=True, before=4, after=4, keep_with_next=True)),
        ("Gate Table Header", WD_STYLE_TYPE.PARAGRAPH, dict(size=9.5, color=DARK_BLUE, bold=True, keep_together=True)),
        ("Gate Table Body", WD_STYLE_TYPE.PARAGRAPH, dict(size=9.5, keep_together=True)),
        ("Policy Number", WD_STYLE_TYPE.PARAGRAPH, dict(size=11, after=8, line_spacing=1.167, keep_together=True)),
    ]
    for name, kind, options in additions:
        configure_paragraph_style(styles.add_style(name, kind), **options)
    for name, options in (
        ("Metadata Label", dict(size=10.5, color=BLACK, bold=True)),
        ("Directive Label", dict(size=11, color=DARK_BLUE, bold=True)),
    ):
        set_style_font(styles.add_style(name, WD_STYLE_TYPE.CHARACTER), name=CALIBRI, **options)
    configure_paragraph_style(styles["Header"], size=8.5, color=MUTED)
    configure_paragraph_style(styles["Footer"], size=8.5, color=MUTED)


def add_numbering_definition(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_id = max([int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))], default=0) + 1
    num_id = max([int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))], default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1."), ("w:lvlJc", "left")):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        level.append(element)
    properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "720")
    indentation.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    properties.extend([tabs, indentation, spacing])
    level.append(properties)
    abstract.append(level)
    numbering.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    numbering.append(number)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    numbers = ensure_child(properties, "w:numPr")
    level = ensure_child(numbers, "w:ilvl")
    level.set(qn("w:val"), "0")
    number = ensure_child(numbers, "w:numId")
    number.set(qn("w:val"), str(num_id))


def add_bookmarked_heading(document: Document, text: str, bookmark_name: str, bookmark_id: int):
    paragraph = document.add_paragraph(style="Heading 1")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    paragraph._p.append(start)
    paragraph.add_run(text)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(end)
    return paragraph


def set_cell_margins(cell) -> None:
    margins = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcMar")
    for side, width in CELL_MARGINS_DXA.items():
        margin = ensure_child(margins, f"w:{side}")
        margin.set(qn("w:w"), str(width))
        margin.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    element = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcW")
    element.set(qn("w:type"), "dxa")
    element.set(qn("w:w"), str(width))
    cell.width = Twips(width)


def set_cell_fill(cell, color: str) -> None:
    shading = ensure_child(cell._tc.get_or_add_tcPr(), "w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)


def set_table_geometry(table) -> None:
    if sum(GATE_COLUMN_WIDTHS_DXA) != CONTENT_WIDTH_DXA:
        raise AssertionError("Gate table columns must sum to 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    for tag, width in (("w:tblW", CONTENT_WIDTH_DXA), ("w:tblInd", TABLE_INDENT_DXA)):
        element = ensure_child(properties, tag)
        element.set(qn("w:type"), "dxa")
        element.set(qn("w:w"), str(width))
    layout = ensure_child(properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    caption = ensure_child(properties, "w:tblCaption")
    caption.set(qn("w:val"), "Eligibility-gates")
    description = ensure_child(properties, "w:tblDescription")
    description.set(qn("w:val"), "Four mandatory supplier eligibility gates")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in GATE_COLUMN_WIDTHS_DXA:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, GATE_COLUMN_WIDTHS_DXA[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    borders = ensure_child(properties, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = ensure_child(borders, f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), BORDER)


def append_field(paragraph, instruction: str) -> None:
    for tag, attributes, text_value in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, instruction),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "1"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        element = OxmlElement(tag)
        for key, value in attributes.items():
            element.set(qn(key), value)
        element.text = text_value
        paragraph.add_run()._r.append(element)


def configure_section(document: Document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    header = section.header.paragraphs[0]
    header.style = document.styles["Header"]
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    header.add_run("Supplier Selection Memo")
    header.add_run("\tCritical service sourcing")
    footer = section.footer.paragraphs[0]
    footer.style = document.styles["Footer"]
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run("Synthetic benchmark fixture")
    footer.add_run("\tPage ")
    append_field(footer, " PAGE ")
    footer.add_run(" of ")
    append_field(footer, " NUMPAGES ")


def add_metadata(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph(style="Masthead Metadata")
    run = paragraph.add_run(f"{label}: ")
    run.style = document.styles["Metadata Label"]
    paragraph.add_run(value)


def set_table_cell_text(cell, text: str, *, header: bool) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.style = "Gate Table Header" if header else "Gate Table Body"
    paragraph.add_run(text)
    if header:
        set_cell_fill(cell, LIGHT_GRAY)


def build_document() -> Document:
    document = Document()
    configure_styles(document)
    configure_section(document)
    numbering_id = add_numbering_definition(document)
    properties = document.core_properties
    properties.author = properties.last_modified_by = "Local LLM Benchmark Lab"
    properties.title = "Supplier Selection Policy"
    properties.subject = "Synthetic multi-format benchmark fixture"
    properties.keywords = "synthetic benchmark, supplier selection, eligibility gates"
    properties.comments = "Project-authored synthetic policy fixture for read-only document analysis."
    fixed_timestamp = datetime(2026, 9, 8, 16, 0, tzinfo=timezone.utc)
    properties.created = properties.modified = fixed_timestamp
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings.element.append(update_fields)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    document.add_paragraph("SUPPLIER SELECTION MEMO", style="Masthead Title")
    document.add_paragraph("Eligibility and ranking policy for the critical service award", style="Masthead Subtitle")
    add_metadata(document, "To", "Sourcing Committee")
    add_metadata(document, "From", "Procurement Governance Office")
    add_metadata(document, "Date", "8 September 2026")
    add_metadata(document, "Re", "Supplier eligibility, primary award and reserve designation")
    add_metadata(document, "Status", "Approved policy for cutoff review")
    rule = document.add_paragraph()
    border = ensure_child(rule._p.get_or_add_pPr(), "w:pBdr")
    bottom = ensure_child(border, "w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "12"), ("w:space", "6"), ("w:color", DARK_BLUE)):
        bottom.set(qn(key), value)
    rule.paragraph_format.space_after = Pt(10)
    directive = document.add_paragraph()
    label = directive.add_run("Directive. ")
    label.style = document.styles["Directive Label"]
    directive.add_run("Test supplier eligibility first. Rank only suppliers that pass every eligibility gate.")

    add_bookmarked_heading(document, "Eligibility gates", "Eligibility_gates", 1)
    document.add_paragraph(
        "A supplier is eligible only with signed-final assurance rating A or B, zero unresolved critical findings, capacity headroom of at least 15%, and adjusted first-year cost no greater than EUR 1,200,000."
    )
    document.add_paragraph("Table 1. Eligibility gates", style="Table Citation")
    rows = [
        ("Gate", "Passing criterion", "Evaluation", "Authoritative evidence"),
        ("Assurance rating", "Signed-final rating is A or B.", "Per supplier", "Signed-final assurance appendix at the cutoff"),
        ("Critical findings", "Unresolved critical finding count equals 0.", "Per supplier", "Signed-final assurance appendix at the cutoff"),
        ("Capacity headroom", "Available / required capacity - 1 is at least 15%.", "Per supplier", "Bid Evaluation, Bid Analysis sheet"),
        ("Adjusted cost", "Adjusted first-year cost is at most EUR 1,200,000.", "Per supplier", "Bid Evaluation, Bid Analysis sheet"),
    ]
    table = document.add_table(rows=len(rows), cols=4)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            set_table_cell_text(table.cell(row_index, column_index), value, header=row_index == 0)
    set_table_geometry(table)
    document.add_paragraph("Policy owner: Procurement Governance Office. Applies to the 2026-09-08 cutoff review.", style="Table Citation")

    selection = add_bookmarked_heading(document, "Selection sequence", "Selection_sequence", 2)
    selection.paragraph_format.page_break_before = True
    document.add_paragraph(
        "First exclude every ineligible supplier. Then select the eligible supplier with the lowest adjusted first-year cost as primary and the next-lowest eligible supplier as reserve. If only one supplier is eligible, make a single award; if none is eligible, make no award."
    )
    document.add_paragraph("Decision procedure", style="Heading 2")
    for item in (
        "Establish the signed-final assurance evidence at the cutoff and recompute adjusted cost and capacity headroom from the bid workbook.",
        "Exclude every supplier that fails any eligibility gate; commercial rank cannot cure a failed gate.",
        "Rank only the remaining eligible suppliers by adjusted first-year cost, lowest first, assigning primary and reserve in that order.",
    ):
        paragraph = document.add_paragraph(item, style="Policy Number")
        apply_numbering(paragraph, numbering_id)

    add_bookmarked_heading(document, "Evidence authority", "Evidence_authority", 3)
    document.add_paragraph(
        "The signed-final assurance appendix at the cutoff supersedes drafts. A commercial dashboard is contextual and cannot waive an assurance or capacity gate, even when it recommends the lowest-cost supplier."
    )
    add_bookmarked_heading(document, "Decision cutoff", "Decision_cutoff", 4)
    document.add_paragraph(
        "The award decision uses evidence finalized by 2026-09-08 16:00 UTC. Later changes belong to a new decision cycle and cannot alter this frozen review retroactively."
    )
    return document


def main() -> None:
    args = parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build_document().save(args.output)


if __name__ == "__main__":
    main()
