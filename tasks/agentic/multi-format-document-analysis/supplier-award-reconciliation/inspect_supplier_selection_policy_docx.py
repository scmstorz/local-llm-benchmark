#!/usr/bin/env python3
"""Run deterministic structural and semantic checks for supplier-selection-policy.docx."""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS}
W = f"{{{W_NS}}}"
EXPECTED_TABLE = [
    ["Gate", "Passing criterion", "Evaluation", "Authoritative evidence"],
    ["Assurance rating", "Signed-final rating is A or B.", "Per supplier", "Signed-final assurance appendix at the cutoff"],
    ["Critical findings", "Unresolved critical finding count equals 0.", "Per supplier", "Signed-final assurance appendix at the cutoff"],
    ["Capacity headroom", "Available / required capacity - 1 is at least 15%.", "Per supplier", "Bid Evaluation, Bid Analysis sheet"],
    ["Adjusted cost", "Adjusted first-year cost is at most EUR 1,200,000.", "Per supplier", "Bid Evaluation, Bid Analysis sheet"],
]
EXPECTED_HEADINGS = ["Eligibility gates", "Selection sequence", "Decision procedure", "Evidence authority", "Decision cutoff"]
EXPECTED_BOOKMARKS = {"Eligibility_gates", "Selection_sequence", "Evidence_authority", "Decision_cutoff"}
EXPECTED_GRID = [1800, 2720, 1760, 3080]
EXPECTED_CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--document", required=True, type=Path)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def integer_attribute(element, name: str) -> int:
    if element is None:
        raise AssertionError(f"Missing element for attribute {name}")
    value = element.get(f"{W}{name}")
    if value is None:
        raise AssertionError(f"Missing w:{name} on {element.tag}")
    return int(value)


def style_element(styles_root, style_id: str):
    matches = styles_root.xpath(f'./w:style[@w:styleId="{style_id}"]', namespaces=NS)
    if len(matches) != 1:
        raise AssertionError(f"Expected one style {style_id}, found {len(matches)}")
    return matches[0]


def assert_style(styles_root, style_id: str, *, half_points: int, before: int, after: int, line: int, color: str) -> None:
    style = style_element(styles_root, style_id)
    fonts = style.find("w:rPr/w:rFonts", namespaces=NS)
    if fonts is None or fonts.get(f"{W}ascii") != "Calibri" or fonts.get(f"{W}hAnsi") != "Calibri":
        raise AssertionError(f"{style_id}: unexpected fonts")
    if integer_attribute(style.find("w:rPr/w:sz", namespaces=NS), "val") != half_points:
        raise AssertionError(f"{style_id}: unexpected font size")
    spacing = style.find("w:pPr/w:spacing", namespaces=NS)
    observed = {key: integer_attribute(spacing, key) for key in ("before", "after", "line")}
    if observed != {"before": before, "after": after, "line": line} or spacing.get(f"{W}lineRule") != "auto":
        raise AssertionError(f"{style_id}: unexpected spacing {observed!r}")
    color_element = style.find("w:rPr/w:color", namespaces=NS)
    if color_element is None or color_element.get(f"{W}val") != color:
        raise AssertionError(f"{style_id}: unexpected color")


def main() -> None:
    args = parse_args()
    document_path = args.document.resolve()
    with zipfile.ZipFile(document_path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"Corrupt ZIP member: {bad_member}")
        members = set(archive.namelist())
        if "word/vbaProject.bin" in members or any(name.startswith("word/comments") for name in members):
            raise AssertionError("Document contains prohibited active or comment content")
        for relationship_path in sorted(name for name in members if name.endswith(".rels")):
            root = etree.fromstring(archive.read(relationship_path))
            if root.xpath('./rel:Relationship[@TargetMode="External"]', namespaces={"rel": PKG_REL_NS}):
                raise AssertionError(f"External relationship in {relationship_path}")
        document_root = etree.fromstring(archive.read("word/document.xml"))
        styles_root = etree.fromstring(archive.read("word/styles.xml"))
        numbering_root = etree.fromstring(archive.read("word/numbering.xml"))
        footer_root = etree.fromstring(archive.read("word/footer1.xml"))

    document = Document(document_path)
    if len(document.sections) != 1:
        raise AssertionError(f"Expected one section, found {len(document.sections)}")
    section = document.sections[0]
    observed_section = {
        "page_width": section.page_width.twips, "page_height": section.page_height.twips,
        "top_margin": section.top_margin.twips, "right_margin": section.right_margin.twips,
        "bottom_margin": section.bottom_margin.twips, "left_margin": section.left_margin.twips,
        "header_distance": section.header_distance.twips, "footer_distance": section.footer_distance.twips,
    }
    expected_section = {
        "page_width": 12240, "page_height": 15840, "top_margin": 1440, "right_margin": 1440,
        "bottom_margin": 1440, "left_margin": 1440, "header_distance": 708, "footer_distance": 708,
    }
    if observed_section != expected_section:
        raise AssertionError(f"Unexpected section geometry: {observed_section!r}")

    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    if headings != EXPECTED_HEADINGS:
        raise AssertionError(f"Unexpected heading sequence: {headings!r}")
    heading_styles = {p.text: p.style.name for p in document.paragraphs if p.style.name.startswith("Heading")}
    if heading_styles["Decision procedure"] != "Heading 2" or any(
        heading_styles[name] != "Heading 1" for name in EXPECTED_HEADINGS if name != "Decision procedure"
    ):
        raise AssertionError(f"Unexpected heading levels: {heading_styles!r}")

    bookmarks = {item.get(f"{W}name") for item in document_root.xpath(".//w:bookmarkStart", namespaces=NS)}
    if not EXPECTED_BOOKMARKS.issubset(bookmarks):
        raise AssertionError(f"Missing bookmarks: {EXPECTED_BOOKMARKS - bookmarks}")
    if len(document.tables) != 1:
        raise AssertionError(f"Expected one table, found {len(document.tables)}")
    table_values = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
    if table_values != EXPECTED_TABLE:
        raise AssertionError(f"Unexpected Eligibility-gates table: {table_values!r}")

    table_xml = document_root.find(".//w:tbl", namespaces=NS)
    table_width = table_xml.find("w:tblPr/w:tblW", namespaces=NS)
    table_indent = table_xml.find("w:tblPr/w:tblInd", namespaces=NS)
    table_caption = table_xml.find("w:tblPr/w:tblCaption", namespaces=NS)
    if integer_attribute(table_width, "w") != 9360 or table_width.get(f"{W}type") != "dxa":
        raise AssertionError("Table width must be 9360 DXA")
    if integer_attribute(table_indent, "w") != 120 or table_indent.get(f"{W}type") != "dxa":
        raise AssertionError("Table indent must be 120 DXA")
    if table_caption is None or table_caption.get(f"{W}val") != "Eligibility-gates":
        raise AssertionError("Table caption anchor must be Eligibility-gates")
    grid = [integer_attribute(column, "w") for column in table_xml.xpath("w:tblGrid/w:gridCol", namespaces=NS)]
    if grid != EXPECTED_GRID:
        raise AssertionError(f"Unexpected table grid: {grid!r}")
    for row_number, row in enumerate(table_xml.xpath("w:tr", namespaces=NS), start=1):
        widths = [integer_attribute(cell.find("w:tcPr/w:tcW", namespaces=NS), "w") for cell in row.xpath("w:tc", namespaces=NS)]
        if widths != EXPECTED_GRID:
            raise AssertionError(f"Row {row_number}: unexpected cell widths {widths!r}")
        for cell in row.xpath("w:tc", namespaces=NS):
            for side, expected in EXPECTED_CELL_MARGINS.items():
                if integer_attribute(cell.find(f"w:tcPr/w:tcMar/w:{side}", namespaces=NS), "w") != expected:
                    raise AssertionError(f"Row {row_number}: unexpected {side} margin")
    fills = [item.get(f"{W}fill") for item in table_xml.xpath("w:tr[1]/w:tc/w:tcPr/w:shd", namespaces=NS)]
    if fills != ["F2F4F7"] * 4:
        raise AssertionError(f"Unexpected header fills: {fills!r}")

    for style_id, values in {
        "Normal": (22, 0, 120, 264, "000000"),
        "Heading1": (32, 320, 160, 240, "2E74B5"),
        "Heading2": (26, 240, 120, 240, "2E74B5"),
        "Heading3": (24, 160, 80, 240, "1F4D78"),
        "TableCitation": (18, 80, 80, 240, "666666"),
        "PolicyNumber": (22, 0, 160, 280, "000000"),
    }.items():
        assert_style(styles_root, style_id, half_points=values[0], before=values[1], after=values[2], line=values[3], color=values[4])

    numbered = document_root.xpath('.//w:p[w:pPr/w:pStyle[@w:val="PolicyNumber"]]', namespaces=NS)
    if len(numbered) != 3:
        raise AssertionError(f"Expected three numbered paragraphs, found {len(numbered)}")
    num_ids = {p.find("w:pPr/w:numPr/w:numId", namespaces=NS).get(f"{W}val") for p in numbered}
    if len(num_ids) != 1:
        raise AssertionError(f"Unexpected numbering IDs: {num_ids!r}")
    num_id = next(iter(num_ids))
    number = numbering_root.xpath(f'./w:num[@w:numId="{num_id}"]', namespaces=NS)[0]
    abstract_id = number.find("w:abstractNumId", namespaces=NS).get(f"{W}val")
    abstract = numbering_root.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]', namespaces=NS)[0]
    indent = abstract.find("w:lvl/w:pPr/w:ind", namespaces=NS)
    tab = abstract.find("w:lvl/w:pPr/w:tabs/w:tab", namespaces=NS)
    spacing = abstract.find("w:lvl/w:pPr/w:spacing", namespaces=NS)
    if (integer_attribute(indent, "left"), integer_attribute(indent, "hanging"), integer_attribute(tab, "pos"), integer_attribute(spacing, "after"), integer_attribute(spacing, "line")) != (720, 360, 720, 160, 280):
        raise AssertionError("Numbering geometry does not match the preset")

    required_fragments = [
        "A supplier is eligible only with signed-final assurance rating A or B",
        "First exclude every ineligible supplier. Then select the eligible supplier with the lowest adjusted first-year cost as primary",
        "The signed-final assurance appendix at the cutoff supersedes drafts.",
        "The award decision uses evidence finalized by 2026-09-08 16:00 UTC.",
    ]
    body_text = "\n".join(p.text for p in document.paragraphs)
    for fragment in required_fragments:
        if fragment not in body_text:
            raise AssertionError(f"Missing required policy fragment: {fragment}")
    footer_fields = {text.strip() for text in footer_root.xpath(".//w:instrText/text()", namespaces=NS)}
    if footer_fields != {"PAGE", "NUMPAGES"}:
        raise AssertionError(f"Unexpected footer fields: {footer_fields!r}")

    result = {
        "status": "passed", "document": str(document_path), "sha256": sha256_file(document_path),
        "size_bytes": document_path.stat().st_size, "sections": 1,
        "section_geometry_twips": observed_section, "headings": headings,
        "bookmarks": sorted(EXPECTED_BOOKMARKS),
        "table": {"caption": "Eligibility-gates", "rows": len(EXPECTED_TABLE), "columns": 4, "width_dxa": 9360, "indent_dxa": 120, "grid_dxa": EXPECTED_GRID, "cell_margins_dxa": EXPECTED_CELL_MARGINS},
        "numbered_paragraphs": 3, "external_relationships": 0, "comments": 0,
        "vba_projects": 0, "policy_fragments": required_fragments,
        "footer_fields": sorted(footer_fields), "author": document.core_properties.author,
    }
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
