#!/usr/bin/env python3
"""Run deterministic structural, formula and cached-value checks for bid-evaluation.xlsx."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import zipfile
from pathlib import Path

import openpyxl
from openpyxl import load_workbook


ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def assert_close(actual: object, expected: float, address: str) -> None:
    if not isinstance(actual, (int, float)) or not math.isclose(actual, expected, rel_tol=0, abs_tol=1e-10):
        raise AssertionError(f"{address}: expected {expected!r}, found {actual!r}")


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    with zipfile.ZipFile(workbook_path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"Corrupt ZIP member: {bad_member}")
        members = set(archive.namelist())
        if any(name.startswith("xl/externalLinks/") for name in members):
            raise AssertionError("Workbook contains an external link")
        if "xl/vbaProject.bin" in members:
            raise AssertionError("Workbook contains a VBA project")

    formula_book = load_workbook(workbook_path, data_only=False)
    value_book = load_workbook(workbook_path, data_only=True)
    if formula_book.sheetnames != ["Dashboard", "Bid Analysis"]:
        raise AssertionError(f"Unexpected worksheets: {formula_book.sheetnames!r}")
    if any(sheet.sheet_state != "visible" for sheet in formula_book.worksheets):
        raise AssertionError("Every worksheet must be visible")

    expected_formulas = {
        "Dashboard!B3": "=INDEX('Bid Analysis'!A5:A8,MATCH(MIN('Bid Analysis'!E5:E8),'Bid Analysis'!E5:E8,0))",
        "Dashboard!B4": "=MIN('Bid Analysis'!E5:E8)",
    }
    for row in range(5, 9):
        expected_formulas.update({
            f"Bid Analysis!E{row}": f"=B{row}+C{row}-D{row}",
            f"Bid Analysis!H{row}": f"=IFERROR(F{row}/G{row}-1,0)",
            f"Bid Analysis!I{row}": f'=IF(E{row}<=$M$3,"PASS","HOLD")',
            f"Bid Analysis!J{row}": f'=IF(H{row}>=$M$4,"PASS","HOLD")',
        })
    observed_formulas: dict[str, str] = {}
    cached_errors: list[str] = []
    for sheet in formula_book.worksheets:
        cached_sheet = value_book[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                address = f"{sheet.title}!{cell.coordinate}"
                if cell.data_type == "f":
                    observed_formulas[address] = cell.value
                if cached_sheet[cell.coordinate].value in ERROR_VALUES:
                    cached_errors.append(address)
    if observed_formulas != expected_formulas:
        raise AssertionError("Formula mismatch:\n" + json.dumps(observed_formulas, indent=2, sort_keys=True))
    if cached_errors:
        raise AssertionError(f"Cached value errors: {cached_errors!r}")

    dashboard = value_book["Dashboard"]
    if dashboard["B3"].value != "Cedar" or dashboard["B5"].value != "Lowest adjusted first-year cost":
        raise AssertionError("Dashboard recommendation anchor is incorrect")
    assert_close(dashboard["B4"].value, 1080000, "Dashboard!B4")
    if [dashboard[f"B{row}"].value for row in (7, 8)] != [
        "Commercial ranking only", "Assurance findings are not incorporated"
    ]:
        raise AssertionError("Dashboard scope anchor is incorrect")

    expected_rows = [
        (5, "Atlas", 1080000, 120000, 60000, 1140000, 116, 100, 0.16, "PASS", "PASS"),
        (6, "Boreal", 1120000, 80000, 35000, 1165000, 125, 100, 0.25, "PASS", "PASS"),
        (7, "Cedar", 1030000, 90000, 40000, 1080000, 120, 100, 0.20, "PASS", "PASS"),
        (8, "Delta", 1020000, 120000, 10000, 1130000, 112, 100, 0.12, "PASS", "HOLD"),
    ]
    analysis = value_book["Bid Analysis"]
    observed_rows = []
    for expected in expected_rows:
        row_number, *values = expected
        actual = [analysis.cell(row_number, column).value for column in range(1, 11)]
        if actual[:8] != values[:8]:
            for index, (observed, wanted) in enumerate(zip(actual[:8], values[:8]), start=1):
                if isinstance(wanted, float):
                    assert_close(observed, wanted, f"Bid Analysis!{openpyxl.utils.get_column_letter(index)}{row_number}")
                elif observed != wanted:
                    raise AssertionError(f"Bid Analysis row {row_number}: {actual!r}")
        if actual[8:] != values[8:]:
            raise AssertionError(f"Bid Analysis row {row_number}: {actual!r}")
        observed_rows.append({"range": f"B{row_number}:J{row_number}", "supplier": values[0], "values": actual[1:]})

    if formula_book["Bid Analysis"].tables["SupplierBidTable"].ref != "A4:J8":
        raise AssertionError("SupplierBidTable must cover A4:J8")
    expected_print_areas = {"Dashboard": "'Dashboard'!$A$1:$F$11", "Bid Analysis": "'Bid Analysis'!$A$1:$M$11"}
    for sheet_name, expected_area in expected_print_areas.items():
        if str(formula_book[sheet_name].print_area) != expected_area:
            raise AssertionError(f"Unexpected {sheet_name} print area: {formula_book[sheet_name].print_area}")
    if any(sheet.page_setup.orientation != "landscape" for sheet in formula_book.worksheets):
        raise AssertionError("Every worksheet must render in landscape orientation")

    result = {
        "status": "passed", "workbook": str(workbook_path),
        "sha256": sha256_file(workbook_path), "size_bytes": workbook_path.stat().st_size,
        "openpyxl_version": openpyxl.__version__, "worksheets": formula_book.sheetnames,
        "formula_count": len(observed_formulas), "formula_error_count": 0,
        "cached_value_error_count": 0,
        "dashboard_anchor": {"range": "B3:B5", "recommended_supplier": dashboard["B3"].value, "lowest_cost": dashboard["B4"].value, "basis": dashboard["B5"].value},
        "dashboard_scope_anchor": {"range": "B7:B8", "values": [dashboard["B7"].value, dashboard["B8"].value]},
        "bid_rows": observed_rows, "formulas": observed_formulas,
        "print_areas": expected_print_areas, "external_links": 0, "vba_projects": 0,
    }
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
